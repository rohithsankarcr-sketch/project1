
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt
import pptx.util as util

# Color palette
BG_DARK      = RGBColor(0x03, 0x06, 0x09)   # near-black
BG_CARD      = RGBColor(0x0d, 0x18, 0x24)   # dark navy
ACCENT_CYAN  = RGBColor(0x00, 0xf0, 0xff)   # AYAS cyan
ACCENT_GOLD  = RGBColor(0xd4, 0xa0, 0x34)   # sacred gold
ACCENT_RED   = RGBColor(0xff, 0x33, 0x44)   # danger red
TEXT_WHITE   = RGBColor(0xff, 0xff, 0xff)
TEXT_LIGHT   = RGBColor(0xc9, 0xd1, 0xd9)
TEXT_DIM     = RGBColor(0x8b, 0x9b, 0xb4)
GREEN_OK     = RGBColor(0x00, 0xe5, 0x76)

SLIDE_W = Inches(13.33)
SLIDE_H = Inches(7.5)

prs = Presentation()
prs.slide_width  = SLIDE_W
prs.slide_height = SLIDE_H

blank = prs.slide_layouts[6]  # completely blank

def add_rect(slide, x, y, w, h, fill_rgb, alpha=None):
    shape = slide.shapes.add_shape(1, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.line.fill.background()
    fill = shape.fill
    fill.solid()
    fill.fore_color.rgb = fill_rgb
    return shape

def add_text(slide, text, x, y, w, h, size, bold=False, color=None, align=PP_ALIGN.LEFT, italic=False, wrap=True):
    txb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf  = txb.text_frame
    tf.word_wrap = wrap
    p   = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color if color else TEXT_WHITE
    return txb

def full_bg(slide, color=BG_DARK):
    add_rect(slide, 0, 0, 13.33, 7.5, color)

def slide_header(slide, title, subtitle=None, title_color=ACCENT_CYAN):
    # Accent bar
    add_rect(slide, 0, 0, 13.33, 0.08, title_color)
    add_text(slide, title, 0.35, 0.18, 12.5, 0.7, 32, bold=True, color=title_color, align=PP_ALIGN.LEFT)
    if subtitle:
        add_text(slide, subtitle, 0.35, 0.82, 12.5, 0.45, 14, color=TEXT_DIM, align=PP_ALIGN.LEFT)

def card(slide, x, y, w, h, color=BG_CARD):
    s = add_rect(slide, x, y, w, h, color)
    return s

# ─────────────────────────── SLIDE 1: TITLE ───────────────────────────
sl = prs.slides.add_slide(blank)
full_bg(sl, BG_DARK)
# Big gradient-like overlay at top
add_rect(sl, 0, 0, 13.33, 3.2, RGBColor(0x04, 0x10, 0x1c))
add_rect(sl, 0, 0, 13.33, 0.12, ACCENT_CYAN)
add_rect(sl, 0, 7.38, 13.33, 0.12, ACCENT_CYAN)

add_text(sl, "AYAS", 0, 1.0, 13.33, 2.0, 96, bold=True, color=ACCENT_CYAN, align=PP_ALIGN.CENTER)
add_text(sl, "T H E   L A S T   W A R D", 0, 2.85, 13.33, 0.6, 22, bold=False, color=ACCENT_GOLD, align=PP_ALIGN.CENTER)
add_text(sl, "3D First-Person Kerala Folklore Survival Horror Game", 0, 3.5, 13.33, 0.5, 16, color=TEXT_DIM, align=PP_ALIGN.CENTER)
add_rect(sl, 3.5, 4.15, 6.33, 0.04, ACCENT_CYAN)
add_text(sl, "HOW TO PLAY — OFFICIAL GAMEPLAY GUIDE", 0, 4.35, 13.33, 0.5, 13, bold=True, color=TEXT_LIGHT, align=PP_ALIGN.CENTER)
add_text(sl, "Southern Kerala, 1994  •  A Remote Monsoon Village  •  One Dark Night", 0, 4.9, 13.33, 0.4, 12, color=TEXT_DIM, align=PP_ALIGN.CENTER)

# Bottom badge
add_rect(sl, 4.5, 6.4, 4.33, 0.65, RGBColor(0x08, 0x1c, 0x2c))
add_text(sl, "AYAS MK-IV  •  ANOMALY DETECTOR  •  FIELD RESEARCH UNIT", 4.5, 6.48, 4.33, 0.5, 9, color=ACCENT_CYAN, align=PP_ALIGN.CENTER)

# ─────────────────────────── SLIDE 2: STORY ───────────────────────────
sl = prs.slides.add_slide(blank)
full_bg(sl)
slide_header(sl, "THE STORY", "What happened in this village?", ACCENT_CYAN)
add_rect(sl, 0, 1.25, 13.33, 0.04, RGBColor(0x1a, 0x33, 0x4d))

story_parts = [
    ("📍 The Setting", "Southern Kerala, 1994. A stormy monsoon night on the abandoned Pala Road.", 0.35, 1.45),
    ("👻 The Mystery", "Villagers report hearing familiar voices in the rain, seeing a pale woman near the ancient Pala tree, and finding damaged iron ward stations. People are disappearing.", 0.35, 2.3),
    ("🔬 Your Role", "You are a paranormal investigator equipped with the prototype AYAS (Anomaly & Iron Integrity Monitor) detector. You enter the village to investigate.", 0.35, 3.45),
    ("⚔️ The Threat", "The Yakshi — a supernatural entity from Kerala folklore — is real. She follows you. She lures you with familiar voices. She hunts in the dark.", 7.0, 1.45),
    ("🎯 Your Mission", "Investigate supernatural activities, restore 3 iron ward stations to form the village's protective perimeter, and survive the final chase to reach the Kavu shrine.", 7.0, 2.5),
    ("🌅 The Stakes", "Complete everything before dawn. If you fail — the Yakshi claims you forever.", 7.0, 3.85),
]

for title, body, bx, by in story_parts:
    card(sl, bx, by, 5.8, 1.55)
    add_text(sl, title, bx+0.15, by+0.1, 5.5, 0.4, 12, bold=True, color=ACCENT_GOLD)
    add_text(sl, body, bx+0.15, by+0.48, 5.5, 1.0, 11, color=TEXT_LIGHT)

add_rect(sl, 0.35, 5.25, 12.63, 1.8, RGBColor(0x08, 0x14, 0x20))
add_text(sl, "\"A remote Kerala village has been experiencing unexplained disappearances. Villagers report hearing familiar voices\nin the rain, seeing a woman near an ancient Pala tree, and finding damaged iron ward stations.\"", 0.5, 5.3, 12.3, 1.65, 12, italic=True, color=ACCENT_CYAN, align=PP_ALIGN.CENTER)

# ─────────────────────────── SLIDE 3: CONTROLS ───────────────────────────
sl = prs.slides.add_slide(blank)
full_bg(sl)
slide_header(sl, "PLAYER CONTROLS", "Master these to survive the night")
add_rect(sl, 0, 1.25, 13.33, 0.04, RGBColor(0x1a, 0x33, 0x4d))

controls = [
    ("W / A / S / D", "Move Forward / Left / Backward / Right", ACCENT_CYAN),
    ("MOUSE", "Look Around (360° First-Person Camera)", ACCENT_CYAN),
    ("SHIFT", "Sprint — Drains Stamina, Attracts Yakshi!", ACCENT_RED),
    ("CTRL", "Crouch — Silent & Reduces Detection Range by 50%", GREEN_OK),
    ("F", "Toggle Flashlight (On/Off) — Increases Yakshi sight range!", ACCENT_GOLD),
    ("Q", "Equip / Stow AYAS Detector Device", ACCENT_CYAN),
    ("R (Hold)", "Calibrate AYAS Detector — Stand Still!", ACCENT_GOLD),
    ("E (Hold)", "Interact / Repair Ward Stations (Hold 2.5 seconds)", GREEN_OK),
    ("E (Tap)", "Inspect Evidence Items / Toggle Hiding Spot", GREEN_OK),
    ("G", "Deploy Sacred Ash — Repels Yakshi (Limited supply!)", ACCENT_RED),
    ("TAB", "Open Research Journal & Evidence Log", TEXT_DIM),
    ("ESC", "Pause Menu / Settings", TEXT_DIM),
]

cols = [controls[:6], controls[6:]]
col_xs = [0.3, 6.9]

for col_i, col in enumerate(cols):
    cx = col_xs[col_i]
    for i, (key, desc, kcolor) in enumerate(col):
        ry = 1.45 + i * 0.97
        add_rect(sl, cx, ry, 6.0, 0.85, RGBColor(0x0d, 0x18, 0x24))
        add_rect(sl, cx, ry, 1.35, 0.85, RGBColor(0x04, 0x10, 0x1c))
        add_text(sl, key, cx+0.05, ry+0.18, 1.25, 0.5, 11, bold=True, color=kcolor, align=PP_ALIGN.CENTER)
        add_text(sl, desc, cx+1.42, ry+0.2, 4.5, 0.5, 11, color=TEXT_LIGHT)

# ─────────────────────────── SLIDE 4: AYAS DETECTOR ───────────────────────────
sl = prs.slides.add_slide(blank)
full_bg(sl)
slide_header(sl, "AYAS MK-IV DETECTOR", "Your primary investigation tool — keep it equipped")
add_rect(sl, 0, 1.25, 13.33, 0.04, ACCENT_CYAN)

# Device mockup panel
add_rect(sl, 0.3, 1.45, 4.0, 5.75, RGBColor(0x04, 0x10, 0x1c))
add_rect(sl, 0.35, 1.5, 3.9, 0.5, RGBColor(0x00, 0x40, 0x55))
add_text(sl, "AYAS MK-IV MONITOR", 0.35, 1.52, 3.9, 0.46, 11, bold=True, color=ACCENT_CYAN, align=PP_ALIGN.CENTER)
add_text(sl, "[ ACTIVE SCAN ]", 0.35, 2.0, 3.9, 0.4, 10, color=GREEN_OK, align=PP_ALIGN.CENTER)
add_text(sl, "EM ANOMALY:  42.5 µT  ████████░░", 0.45, 2.5, 3.7, 0.4, 10, color=ACCENT_CYAN)
add_text(sl, "COLD-SPOT:  28.5°C  ↓ DROPPING", 0.45, 3.0, 3.7, 0.4, 10, color=ACCENT_CYAN)
add_text(sl, "VOICE SPEC:  ANOMALOUS  ⚠", 0.45, 3.5, 3.7, 0.4, 10, color=ACCENT_RED)
add_text(sl, "WARD NAV:  WARD A | 24m ↑", 0.45, 4.0, 3.7, 0.4, 10, color=ACCENT_GOLD)
add_rect(sl, 0.45, 4.55, 3.7, 0.6, RGBColor(0x00, 0x18, 0x22))
add_text(sl, "〜〜〜〜∧∧〜〜〜〜∨∨〜〜〜", 0.45, 4.62, 3.7, 0.45, 10, color=GREEN_OK, align=PP_ALIGN.CENTER)
add_text(sl, "SENSORS: ONLINE  |  HOLD [R] TO CALIBRATE", 0.45, 5.25, 3.7, 0.4, 9, color=TEXT_DIM, align=PP_ALIGN.CENTER)

# Four modes
modes = [
    ("📡 EM Anomaly Meter", "Detects electromagnetic disturbances caused by the Yakshi. Needle rises as she gets closer. Watch for sudden spikes — she's near!"),
    ("🌡️ Cold-Spot Temperature", "Ambient temperature drops when the Yakshi is present. Below 20°C = danger zone. Below 10°C = RUN."),
    ("🎙️ Voice Analyzer", "Analyzes nearby voices for paranormal patterns. If meter reads 'ANOMALOUS', the voice you hear is the Yakshi LURING you — do NOT follow it."),
    ("🧭 Ward Compass", "Points toward the nearest unrepaired Iron Ward station. Follow the direction arrow to locate your next objective."),
]

for i, (title, body) in enumerate(modes):
    my = 1.45 + i * 1.45
    card(sl, 4.65, my, 8.35, 1.35)
    add_text(sl, title, 4.8, my+0.08, 8.0, 0.45, 13, bold=True, color=ACCENT_GOLD)
    add_text(sl, body, 4.8, my+0.52, 8.0, 0.75, 11, color=TEXT_LIGHT)

add_rect(sl, 4.65, 7.1, 8.35, 0.04, ACCENT_CYAN)
add_text(sl, "Press [Q] to equip/stow  •  Press [R] while standing still to calibrate  •  Watch the waveform for ghost activity", 4.65, 7.18, 8.35, 0.28, 9, color=TEXT_DIM, align=PP_ALIGN.CENTER)

# ─────────────────────────── SLIDE 5: OBJECTIVES ───────────────────────────
sl = prs.slides.add_slide(blank)
full_bg(sl)
slide_header(sl, "8 MISSION OBJECTIVES", "Complete them in order to escape before dawn", ACCENT_GOLD)
add_rect(sl, 0, 1.25, 13.33, 0.04, ACCENT_GOLD)

objectives = [
    ("1", "Enter the Village", "Walk the Pala Road to reach the abandoned ancestral house. Learn the controls.", ACCENT_CYAN, "EXPLORATION"),
    ("2", "Calibrate AYAS Detector", "Find a baseline zone, stand still, and hold [R] to calibrate your device.", ACCENT_CYAN, "DETECTIVE"),
    ("3", "Investigate the Pala Tree", "Approach the ancient Pala tree. Document paranormal anomalies.", ACCENT_GOLD, "INVESTIGATE"),
    ("4", "Restore Ward A", "Locate Iron Ward Post A on the Village Path. Hold [E] to repair.", GREEN_OK, "REPAIR"),
    ("5", "Survive & Restore Ward B", "Ward B is near the Pala Tree. The Yakshi is active here — be stealthy!", ACCENT_RED, "DANGER"),
    ("6", "Recover Evidence", "Find 3 case files & records inside the ancestral house. Press [E] to collect.", ACCENT_CYAN, "DETECTIVE"),
    ("7", "Restore Ward C", "Final perimeter breach at the ancestral house. Repair before she breaks through!", ACCENT_RED, "CRITICAL"),
    ("8", "Escape Before Dawn", "Sprint to the sacred shrine! Scatter Sacred Ash if she chases you!", ACCENT_GOLD, "ESCAPE"),
]

for i, (num, title, desc, col, tag) in enumerate(objectives):
    col_i = i % 2
    row_i = i // 2
    ox = 0.3 + col_i * 6.6
    oy = 1.42 + row_i * 1.48
    card(sl, ox, oy, 6.3, 1.38)
    add_rect(sl, ox, oy, 0.55, 1.38, RGBColor(0x04, 0x10, 0x1c))
    add_text(sl, num, ox, oy+0.42, 0.55, 0.55, 20, bold=True, color=col, align=PP_ALIGN.CENTER)
    add_rect(sl, ox+5.5, oy+0.08, 0.72, 0.35, RGBColor(0x04, 0x10, 0x1c))
    add_text(sl, tag, ox+5.5, oy+0.1, 0.72, 0.35, 7, bold=True, color=col, align=PP_ALIGN.CENTER)
    add_text(sl, title, ox+0.65, oy+0.08, 4.7, 0.45, 12, bold=True, color=col)
    add_text(sl, desc, ox+0.65, oy+0.52, 4.85, 0.8, 10, color=TEXT_LIGHT)

# ─────────────────────────── SLIDE 6: THE YAKSHI ───────────────────────────
sl = prs.slides.add_slide(blank)
full_bg(sl)
add_rect(sl, 0, 0, 13.33, 7.5, RGBColor(0x06, 0x02, 0x04))
slide_header(sl, "THE YAKSHI — KNOW YOUR ENEMY", "She is always watching. She is always listening.", ACCENT_RED)
add_rect(sl, 0, 1.25, 13.33, 0.04, ACCENT_RED)

# Entity description
add_rect(sl, 0.3, 1.45, 5.5, 2.8, RGBColor(0x10, 0x04, 0x06))
add_text(sl, "Who Is She?", 0.45, 1.52, 5.2, 0.45, 14, bold=True, color=ACCENT_RED)
add_text(sl, "The Yakshi is a supernatural entity from Kerala folklore. She appears as a beautiful woman in a white Kasavu saree with long dark hair. She lures victims with sweet voices and familiar calls before striking.\n\nShe cannot cross a fully repaired iron ward perimeter — that is your only protection.", 0.45, 1.95, 5.2, 2.2, 11, color=TEXT_LIGHT)

# 9 AI States
add_rect(sl, 6.1, 1.45, 6.9, 5.75, RGBColor(0x0d, 0x08, 0x0c))
add_text(sl, "⚡ Her 9 Behavioral States", 6.25, 1.5, 6.6, 0.45, 13, bold=True, color=ACCENT_RED)

states = [
    ("💤 Dormant", "Inactive — waiting to be triggered", TEXT_DIM),
    ("👁️ Stalking", "Circling at distance (14–20m), watching you", ACCENT_RED),
    ("🎭 Luring", "Calling your name with a familiar voice — DON'T follow!", ACCENT_RED),
    ("🔍 Investigating", "Heard a noise — moving toward the sound", ACCENT_GOLD),
    ("🏃 Hunting", "Line of sight confirmed — she charges at full speed!", ACCENT_RED),
    ("🔎 Searching", "Lost you — checking your last known position", ACCENT_GOLD),
    ("💨 Disrupted", "Hit by Sacred Ash or Ward energy — fleeing!", GREEN_OK),
    ("📍 Repositioning", "Moving to a new flank position — she's adaptive!", ACCENT_GOLD),
    ("💀 Attack", "Reached you — GAME OVER", ACCENT_RED),
]
for i, (state, desc, sc) in enumerate(states):
    sy = 1.95 + i * 0.57
    add_text(sl, state, 6.2, sy, 2.2, 0.52, 9, bold=True, color=sc)
    add_text(sl, desc, 8.35, sy, 4.55, 0.52, 9, color=TEXT_LIGHT)

# Behavior tips
behaviors = [
    ("🔇 She HEARS you sprint", "Sprinting within 30m alerts her. Walk or crouch near her.", ACCENT_RED),
    ("👁️ She SEES your light", "Your flashlight increases her sight range. Turn it off to hide.", ACCENT_GOLD),
    ("👂 Voice Lures are TRAPS", "If AYAS says ANOMALOUS, the voice you hear is her. Don't go toward it.", ACCENT_RED),
    ("🛡️ Iron Wards repel her", "A repaired ward station repels her if she's within 24 meters.", GREEN_OK),
]
for i, (title, body, bc) in enumerate(behaviors):
    bx = 0.3 + (i % 2) * 2.85
    by = 4.45 + (i // 2) * 1.25
    add_rect(sl, bx, by, 2.75, 1.1, RGBColor(0x14, 0x06, 0x08))
    add_text(sl, title, bx+0.1, by+0.08, 2.55, 0.4, 10, bold=True, color=bc)
    add_text(sl, body, bx+0.1, by+0.48, 2.55, 0.57, 9, color=TEXT_LIGHT)

# ─────────────────────────── SLIDE 7: SURVIVAL TIPS ───────────────────────────
sl = prs.slides.add_slide(blank)
full_bg(sl)
slide_header(sl, "SURVIVAL MECHANICS", "Use every tool you have — or die trying", GREEN_OK)
add_rect(sl, 0, 1.25, 13.33, 0.04, GREEN_OK)

mechanics = [
    ("🦶 Crouching", "[CTRL]", "Crouch to reduce footstep noise and halve detection range. Yakshi's sight range drops by 50% when you crouch. Use this near her!", GREEN_OK),
    ("🏃 Stamina System", "[SHIFT]", "Sprinting drains stamina. When exhausted, you can only walk. Rest in a hidden spot to recover faster. Never sprint near the Yakshi!", ACCENT_GOLD),
    ("🙈 Hiding Spots", "[E] near bin", "Find the wooden rice bin (Nellara) on the house veranda. Hide inside when hunted — she CANNOT see you while you hide.", GREEN_OK),
    ("🌿 Sacred Ash", "[G] key", "You have 2 pouches of Sacred Ash (Bhasmam). Deploying it creates a glowing barrier and repels the Yakshi if she's within 25 meters. Use wisely!", ACCENT_RED),
    ("💡 Flashlight", "[F] key", "Your flashlight is essential to see. BUT it increases the Yakshi's sight range. In stealth sections, consider turning it off to move unseen.", ACCENT_GOLD),
    ("🗺️ Minimap", "[Always visible]", "Top-right minimap shows your position (blue dot), all 3 ward stations, and the shrine. The Yakshi appears as a red dot when your AYAS detects her.", ACCENT_CYAN),
]

for i, (title, key, body, col) in enumerate(mechanics):
    ci = i % 2
    ri = i // 2
    mx = 0.3 + ci * 6.55
    my = 1.45 + ri * 1.9
    card(sl, mx, my, 6.2, 1.8)
    add_rect(sl, mx, my, 6.2, 0.5, RGBColor(0x04, 0x10, 0x1c))
    add_text(sl, title, mx+0.15, my+0.06, 4.0, 0.4, 13, bold=True, color=col)
    add_rect(sl, mx+4.5, my+0.08, 1.55, 0.35, col)
    add_text(sl, key, mx+4.5, my+0.1, 1.55, 0.32, 10, bold=True, color=BG_DARK, align=PP_ALIGN.CENTER)
    add_text(sl, body, mx+0.15, my+0.6, 5.9, 1.1, 11, color=TEXT_LIGHT)

# ─────────────────────────── SLIDE 8: MINIMAP & HUD ───────────────────────────
sl = prs.slides.add_slide(blank)
full_bg(sl)
slide_header(sl, "HUD & MINIMAP GUIDE", "Understanding your on-screen information")
add_rect(sl, 0, 1.25, 13.33, 0.04, ACCENT_CYAN)

hud_elements = [
    ("📌 TOP LEFT — Objective Card", "Shows your current active objective title, description, and hint. Updates automatically when you complete each objective.", 0.3, 1.45),
    ("🔴 TOP CENTER — Ward Status", "Three dots (WARD A, B, C). Red = damaged. Cyan = repaired. All three cyan = perimeter restored!", 0.3, 2.65),
    ("➕ CENTER — Crosshair & Prompts", "A dot crosshair appears in the center. When near an interactable object, a context prompt appears below it.", 0.3, 3.85),
    ("📊 BOTTOM LEFT — Stamina Bar", "Green bar shows current stamina. Turns red when exhausted. Sacred Ash count shown next to it.", 0.3, 5.05),
    ("🗺️ TOP RIGHT — Minimap", "Live overhead minimap. Blue dot = you. Orange triangles = wards. Green zone = shrine. Red blip = Yakshi (when detected).", 6.9, 1.45),
    ("📡 BOTTOM RIGHT — AYAS Panel", "Live AYAS device readings: EM level, temperature, voice analysis, and ward navigation compass.", 6.9, 2.65),
    ("💀 DEATH SCREEN", "If the Yakshi catches you, the screen shows cause of death. You can retry from checkpoint.", 6.9, 3.85),
    ("🏆 VICTORY SCREEN", "Reach the Kavu shrine with all 3 wards repaired to trigger the victory cinematic and mission summary.", 6.9, 5.05),
]

for title, body, bx, by in hud_elements:
    card(sl, bx, by, 6.2, 1.05)
    add_text(sl, title, bx+0.15, by+0.08, 5.9, 0.38, 11, bold=True, color=ACCENT_CYAN)
    add_text(sl, body, bx+0.15, by+0.48, 5.9, 0.52, 10, color=TEXT_LIGHT)

# ─────────────────────────── SLIDE 9: TIPS & TRICKS ───────────────────────────
sl = prs.slides.add_slide(blank)
full_bg(sl)
slide_header(sl, "PRO TIPS & TRICKS", "Expert survival strategies for the Pala Road", ACCENT_GOLD)
add_rect(sl, 0, 1.25, 13.33, 0.04, ACCENT_GOLD)

tips = [
    ("🎵 Listen to the Rain", "The soundscape tells you where the Yakshi is. Anklet bells (payal sounds) mean she's within 28 meters. Closer chimes = closer danger.", ACCENT_GOLD),
    ("📱 AYAS Never Lies", "If EM is above 80 µT and temperature is dropping, she's very near. Equip your AYAS constantly — it's your best early warning system.", ACCENT_CYAN),
    ("🌙 Manage Your Flashlight", "In open areas, turn the flashlight OFF when you suspect she's near. The darkness is your friend for stealth. Turn it ON only to navigate or read.", ACCENT_GOLD),
    ("⚡ Repair Wards in Order", "The objectives are designed to escalate tension. Don't run ahead — complete each ward in sequence to trigger proper story beats.", GREEN_OK),
    ("🏃 The Final Chase", "Objective 8 begins the climax chase. Sprint toward the shrine. If pursued, scatter Sacred Ash behind you and keep running. DO NOT stop.", ACCENT_RED),
    ("👂 Don't Trust the Voices", "If you hear someone calling your name or a familiar voice — check your AYAS Voice Analyzer first. If it says ANOMALOUS, it's her.", ACCENT_RED),
    ("🌿 Ash as Last Resort", "Sacred Ash pouches are precious — only 2 available. Save them for the final escape chase, not for early encounters where hiding works better.", ACCENT_GOLD),
    ("🔇 Crouch Near the Tree", "Ward B repair is the most dangerous. Crouch constantly near the Pala Tree area. Never sprint. Use the AYAS to track her position.", GREEN_OK),
]

for i, (title, body, col) in enumerate(tips):
    ci = i % 2
    ri = i // 2
    tx = 0.3 + ci * 6.55
    ty = 1.42 + ri * 1.48
    card(sl, tx, ty, 6.2, 1.38)
    add_rect(sl, tx, ty, 6.2, 0.06, col)
    add_text(sl, title, tx+0.15, ty+0.12, 5.9, 0.42, 12, bold=True, color=col)
    add_text(sl, body, tx+0.15, ty+0.56, 5.9, 0.78, 10, color=TEXT_LIGHT)

# ─────────────────────────── SLIDE 10: QUICK REFERENCE ───────────────────────────
sl = prs.slides.add_slide(blank)
full_bg(sl)
add_rect(sl, 0, 0, 13.33, 7.5, RGBColor(0x03, 0x08, 0x0e))
slide_header(sl, "QUICK REFERENCE SHEET", "Keep this handy while playing", ACCENT_CYAN)
add_rect(sl, 0, 1.25, 13.33, 0.04, ACCENT_CYAN)

# Controls cheat-sheet
add_rect(sl, 0.3, 1.38, 6.0, 5.8, RGBColor(0x05, 0x12, 0x1c))
add_text(sl, "⌨️ CONTROLS AT A GLANCE", 0.45, 1.43, 5.7, 0.45, 13, bold=True, color=ACCENT_CYAN)
quick_ctrl = [
    ("WASD", "Move"),("Mouse", "Look"),("Shift", "Sprint"),("Ctrl", "Crouch"),
    ("F", "Flashlight"),("Q", "AYAS Device"),("R (Hold)", "Calibrate"),
    ("E (Hold)", "Repair Ward"),("E (Tap)", "Collect / Hide"),
    ("G", "Sacred Ash"),("Tab", "Journal"),("Esc", "Pause"),
]
for i, (k, v) in enumerate(quick_ctrl):
    ky = 1.88 + i * 0.41
    add_rect(sl, 0.35, ky, 1.4, 0.38, RGBColor(0x02, 0x08, 0x14))
    add_text(sl, k, 0.35, ky+0.04, 1.4, 0.35, 10, bold=True, color=ACCENT_CYAN, align=PP_ALIGN.CENTER)
    add_text(sl, v, 1.8, ky+0.07, 4.4, 0.32, 10, color=TEXT_LIGHT)

# Danger cheatsheet
add_rect(sl, 6.65, 1.38, 6.35, 2.65, RGBColor(0x12, 0x04, 0x06))
add_text(sl, "☠️ YAKSHI WARNING SIGNS", 6.8, 1.43, 6.0, 0.45, 13, bold=True, color=ACCENT_RED)
warns = [
    "🔊 Anklet bells (payal) = she's within 28m",
    "🌡️ Cold & fog = she's within 15m",
    "⚡ AYAS EM > 80µT = she's near",
    "🎵 Voice calls = she's LURING you (ignore!)",
    "💡 Flashlight flickers = she's transitioning to Hunt",
    "🏃 Rapid footsteps ahead = she's Investigating",
]
for i, w in enumerate(warns):
    add_text(sl, w, 6.8, 1.92 + i * 0.37, 6.0, 0.35, 10, color=ACCENT_RED if i < 3 else ACCENT_GOLD)

# Objective quick-list
add_rect(sl, 6.65, 4.12, 6.35, 3.06, RGBColor(0x04, 0x10, 0x1c))
add_text(sl, "🎯 OBJECTIVE CHECKLIST", 6.8, 4.18, 6.0, 0.45, 13, bold=True, color=ACCENT_GOLD)
obj_quick = [
    "1. Enter Village → reach ancestral house",
    "2. Calibrate AYAS → stand still, hold [R]",
    "3. Investigate Pala Tree → approach & observe",
    "4. Repair Ward A → Village Path, hold [E]",
    "5. Repair Ward B → Pala Tree (DANGER ZONE!)",
    "6. Collect 3 Evidence Items → ancestral house",
    "7. Repair Ward C → house perimeter",
    "8. Escape to shrine → SPRINT & scatter ash!",
]
for i, o in enumerate(obj_quick):
    col_o = GREEN_OK if i < 7 else ACCENT_RED
    add_text(sl, o, 6.8, 4.65 + i * 0.32, 6.0, 0.3, 9, color=col_o)

# ─────────────────────────── SLIDE 11: END SCREEN ───────────────────────────
sl = prs.slides.add_slide(blank)
full_bg(sl, BG_DARK)
add_rect(sl, 0, 0, 13.33, 0.12, ACCENT_CYAN)
add_rect(sl, 0, 7.38, 13.33, 0.12, ACCENT_CYAN)
add_rect(sl, 0, 0, 13.33, 7.5, RGBColor(0x03, 0x06, 0x09))

add_text(sl, "GOOD LUCK, INVESTIGATOR", 0, 1.8, 13.33, 1.0, 36, bold=True, color=ACCENT_CYAN, align=PP_ALIGN.CENTER)
add_text(sl, "The village needs you. The iron wards must hold.", 0, 2.85, 13.33, 0.5, 18, color=ACCENT_GOLD, align=PP_ALIGN.CENTER)
add_rect(sl, 3.5, 3.55, 6.33, 0.05, ACCENT_CYAN)
add_text(sl, "Remember:", 0, 3.75, 13.33, 0.5, 14, bold=True, color=TEXT_WHITE, align=PP_ALIGN.CENTER)
tips_final = [
    "• Crouch often  •  Trust the AYAS  •  Don't follow the voices  •  Repair all 3 wards  •  Save your Sacred Ash"
]
for t in tips_final:
    add_text(sl, t, 1.0, 4.25, 11.33, 0.55, 13, color=TEXT_LIGHT, align=PP_ALIGN.CENTER)

add_text(sl, "\"She cannot permeate conductive iron lines energized with sacred ash.\"", 0, 5.1, 13.33, 0.5, 13, italic=True, color=TEXT_DIM, align=PP_ALIGN.CENTER)
add_text(sl, "— Ritual Copper Plate Record, found in the Ancestral House", 0, 5.55, 13.33, 0.4, 11, italic=True, color=TEXT_DIM, align=PP_ALIGN.CENTER)

add_rect(sl, 4.0, 6.3, 5.33, 0.72, RGBColor(0x06, 0x14, 0x22))
add_text(sl, "AYAS: THE LAST WARD  •  © 2024  •  Kerala Folklore Horror", 4.0, 6.42, 5.33, 0.5, 10, color=ACCENT_CYAN, align=PP_ALIGN.CENTER)

# Save
out_path = r"c:\Users\LOQ\Desktop\bhooo\AYAS_Gameplay_Guide.pptx"
prs.save(out_path)
print(f"Saved: {out_path}")
