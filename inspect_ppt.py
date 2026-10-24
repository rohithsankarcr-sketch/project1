from pptx import Presentation
from pptx.util import Emu

for name in ['AYAS_Gameplay_Guide.pptx']:
    prs = Presentation(name)
    print('FILE', name, 'slides', len(prs.slides), 'size', prs.slide_width, prs.slide_height)
    for i, slide in enumerate(prs.slides):
        print('--- SLIDE', i+1, 'layout', slide.slide_layout.name)
        for sh in slide.shapes:
            info = [sh.shape_type, sh.name]
            try:
                info.append(f'pos=({Emu(sh.left).inches:.2f},{Emu(sh.top).inches:.2f}) size=({Emu(sh.width).inches:.2f}x{Emu(sh.height).inches:.2f})')
            except Exception:
                pass
            print('  SHAPE', info)
            if sh.has_text_frame:
                for p in sh.text_frame.paragraphs:
                    t = ''.join(r.text for r in p.runs)
                    if t.strip():
                        sizes = [r.font.size.pt if r.font.size else None for r in p.runs]
                        colors = []
                        for r in p.runs:
                            try:
                                colors.append(str(r.font.color.rgb))
                            except Exception:
                                colors.append(None)
                        fonts = [r.font.name for r in p.runs]
                        print('     TXT:', repr(t[:120]), 'sizes', sizes, 'colors', colors, 'fonts', fonts)
            if sh.has_table:
                print('     TABLE', len(sh.table.rows), 'x', len(sh.table.columns))
