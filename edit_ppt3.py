from pptx import Presentation
from pptx.util import Pt

prs = Presentation('AYAS_Gameplay_Guide_new.pptx')
slide_layout = prs.slide_layouts[1]
slide = prs.slides.add_slide(slide_layout)
slide.shapes.title.text = "Advanced Features"
body = slide.shapes.placeholders[1].text_frame
body.clear()
for t in ["Adaptive AI, Sound Investigation, Voice Mimicry", "Fake Sightings & Luring", "Dynamic Weather (Rain/Thunder/Lightning)", "Evidence Board with Connections", "Multiple Endings A/B/C", "Relic Weapon [V]", "Health System, Concussion", "Tutorial + 5 Levels"]:
    p = body.add_paragraph()
    p.text = "• " + t
    p.font.size = Pt(17)
prs.save('AYAS_Gameplay_Guide_edited.pptx')
print('ok')
