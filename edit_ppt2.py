from pptx import Presentation
from pptx.util import Pt

prs = Presentation('AYAS_Gameplay_Guide_new.pptx')
slide_layout = prs.slide_layouts[1]
slide = prs.slides.add_slide(slide_layout)
slide.shapes.title.text = "Advanced Features - Latest Update"
body = slide.shapes.placeholders[1].text_frame
body.clear()
items = [
    "Adaptive Yakshi AI - learns patterns",
    "Sound Investigation & Voice Mimicry",
    "Fake Sightings & Luring",
    "Weather: Rain/Thunder/Lightning/Fog",
    "Evidence Board with Clue Connections",
    "Multiple Endings A/B/C",
    "Relic Weapon [V] to Disrupt",
    "Health/Damage System with Restart",
    "Concussion & Aggressive AI",
    "6 Levels (Tutorial + 5)"
]
for t in items:
    p = body.add_paragraph()
    p.text = "• " + t
    p.font.size = Pt(16)

slide = prs.slides.add_slide(slide_layout)
slide.shapes.title.text = "Endings & Survival"
body = slide.shapes.placeholders[1].text_frame
body.clear()
items = ["A: Escape - Reach shrine", "B: The Truth - 2+ evidence", "C: The Last Ward - All 3 wards + 4+ evidence", "Sacred Ash [G] - Limited defense", "Relic [V] - Disrupt ghost", "Hide in cupboards when hunted"]
for t in items:
    p = body.add_paragraph()
    p.text = "• " + t
    p.font.size = Pt(18)

prs.save('AYAS_Gameplay_Guide.pptx')
print('saved')
