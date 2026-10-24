from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
import os

prs = Presentation()
prs.slide_width = Inches(10)
prs.slide_height = Inches(7.5)

# Title
slide = prs.slides.add_slide(prs.slide_layouts[0])
slide.shapes.title.text = "AYAS: THE LAST WARD"
slide.placeholders[1].text = "Complete Gameplay Guide - Updated\nAI, Investigation, Endings, Weather & More"

# Game Overview
slide = prs.slides.add_slide(prs.slide_layouts[1])
slide.shapes.title.text = "Game Overview"
body = slide.shapes.placeholders[1].text_frame
body.clear()
p = body.paragraphs[0]
p.text = "Genre: First-Person Survival Horror (Kerala Folklore)"
p = body.add_paragraph()
p.text = "Setting: Monsoon night, Tharavadu house, Pala tree, Kavu shrine"
p = body.add_paragraph()
p.text = "Antagonist: Intelligent Yakshi with advanced AI"
p = body.add_paragraph()
p.text = "Objective: Investigate, collect evidence, repair 3 wards, reach Kavu shrine"
p = body.add_paragraph()
p.text = "Levels: Tutorial + 5 levels with increasing difficulty"
p = body.add_paragraph()
p.text = "Multiple Endings: A (Escape), B (The Truth), C (The Last Ward)"

# Controls
slide = prs.slides.add_slide(prs.slide_layouts[1])
slide.shapes.title.text = "Controls"
body = slide.shapes.placeholders[1].text_frame
body.clear()
for k,v in [("WASD","Move"),("Mouse","Look"),("Shift","Sprint"),("Ctrl","Crouch"),("E","Interact/Hide/Repair"),("F","Flashlight"),("Q","AYAS Detector"),("R","Calibrate"),("G","Sacred Ash"),("V","Relic Weapon"),("Tab","Journal/Evidence Board"),("Esc","Pause")]:
    p = body.add_paragraph(); p.text = f"{k}: {v}"

# Yakshi AI
slide = prs.slides.add_slide(prs.slide_layouts[1])
slide.shapes.title.text = "Yakshi AI - Advanced Behavior"
body = slide.shapes.placeholders[1].text_frame
body.clear()
for t in ["Adaptive - learns player patterns (sprint/hide/flashlight)", "Sound Investigation - hears footsteps, repairs, interactions", "Voice Mimicry & Luring - false calls to mislead", "Fake Sightings - shadows/silhouettes", "Hunt/Chase - fast, aggressive, can kill (45 damage)", "Limited pursuit time, investigates last known location", "Disrupted by Sacred Ash/Relic", "Searches after losing sight"]:
    p = body.add_paragraph(); p.text = "• " + t

# Survival
slide = prs.slides.add_slide(prs.slide_layouts[1])
slide.shapes.title.text = "Survival Systems"
body = slide.shapes.placeholders[1].text_frame
body.clear()
for t in ["Health Bar - 100 HP, damaged by Yakshi", "Sacred Ash (2 charges) - disrupts Yakshi temporarily", "Relic Weapon [V] - powerful disruption tool", "Hiding - cupboards/containers, enemy checks plausibly", "Crouch - reduces detection", "Stamina - sprint drains, regenerates", "Concussion/Knockback on attacks"]:
    p = body.add_paragraph(); p.text = "• " + t

# Investigation & Evidence
slide = prs.slides.add_slide(prs.slide_layouts[1])
slide.shapes.title.text = "Investigation & Evidence Board"
body = slide.shapes.placeholders[1].text_frame
body.clear()
for t in ["Collect diary pages, photos, reports, notes", "Evidence Board tracks clues with connections", "Journal (Tab) logs all evidence", "Clues reveal Yakshi behavior & ward repairs", "Some clues contradict for mystery"]:
    p = body.add_paragraph(); p.text = "• " + t

# Wards & Puzzles
slide = prs.slides.add_slide(prs.slide_layouts[1])
slide.shapes.title.text = "Wards & Puzzles"
body = slide.shapes.placeholders[1].text_frame
body.clear()
for t in ["3 Iron Wards (A, B, C) to repair", "Each ward has unique repair mechanic", "Repair takes time - Yakshi investigates noise", "AYAS Detector helps locate & analyze", "Restoring wards repels Yakshi"]:
    p = body.add_paragraph(); p.text = "• " + t

# Weather & Atmosphere
slide = prs.slides.add_slide(prs.slide_layouts[1])
slide.shapes.title.text = "Weather & Atmosphere"
body = slide.shapes.placeholders[1].text_frame
body.clear()
for t in ["Dynamic Weather - rain, thunder, lightning, fog", "Spatial Audio - anklets, whispers, footsteps", "Dark lighting, fog, supernatural effects", "Lightning reveals silhouettes", "Rain masks footsteps"]:
    p = body.add_paragraph(); p.text = "• " + t

# Endings
slide = prs.slides.add_slide(prs.slide_layouts[1])
slide.shapes.title.text = "Multiple Endings"
body = slide.shapes.placeholders[1].text_frame
body.clear()
for t in ["Ending A - Escape: Reach shrine with minimal evidence", "Ending B - The Truth: Gather critical evidence (2+)", "Ending C - The Last Ward: Repair all 3 wards + collect evidence (4+)"]:
    p = body.add_paragraph(); p.text = "• " + t

# Tips
slide = prs.slides.add_slide(prs.slide_layouts[1])
slide.shapes.title.text = "Tips to Survive"
body = slide.shapes.placeholders[1].text_frame
body.clear()
for t in ["Use crouch to move quietly", "Listen for anklets & whispers", "Don't follow false voices", "Hide when hunting starts", "Use Sacred Ash as emergency defense", "Use Relic [V] when overwhelmed", "Watch AYAS Detector readings", "Repair wards to create safe zones", "Check Journal for clues"]:
    p = body.add_paragraph(); p.text = "• " + t

# Final
slide = prs.slides.add_slide(prs.slide_layouts[0])
slide.shapes.title.text = "Good Luck"
slide.placeholders[1].text = "AYAS: THE LAST WARD\nSurvive the Night"

out = 'AYAS_Gameplay_Guide_Updated.pptx'
prs.save(out)
print('saved', out)
