from pptx import Presentation
from pptx.util import Pt

# Try to open existing
try:
    prs = Presentation('AYAS_Gameplay_Guide.pptx')
    print('opened')
except:
    from pptx import Presentation
    prs = Presentation()
    print('new')

# Add/update slides with new content
# Add new slide for Advanced Features
from pptx.util import Inches
slide_layout = prs.slide_layouts[1]
slide = prs.slides.add_slide(slide_layout)
slide.shapes.title.text = "NEW: Advanced AI & Features"
body = slide.shapes.placeholders[1].text_frame
body.clear()
items = [
    "Adaptive Yakshi AI - learns sprinting/hiding/flashlight use",
    "Sound Investigation - responds to footsteps, repairs, noise",
    "Voice Mimicry & False Lures - deceptive whispers",
    "Fake Sightings - shadows & silhouettes",
    "Weather System - Rain, Thunder, Lightning, Fog",
    "Evidence Board - Connect clues, form theories",
    "Multiple Endings (A/B/C) based on evidence & wards",
    "Relic Weapon [V] to disrupt ghost objects",
    "Health System with death & restart",
    "Concussion effects on attacks"
]
for t in items:
    p = body.add_paragraph()
    p.text = "• " + t
    p.font.size = Pt(18)

# Add endings slide
slide = prs.slides.add_slide(slide_layout)
slide.shapes.title.text = "Multiple Endings"
body = slide.shapes.placeholders[1].text_frame
body.clear()
for t in ["A - Escape: Reach shrine, unanswered questions", "B - The Truth: 2+ evidence pieces collected", "C - The Last Ward: All 3 wards repaired + 4+ evidence"]:
    p = body.add_paragraph()
    p.text = t
    p.font.size = Pt(20)

prs.save('AYAS_Gameplay_Guide.pptx')
print('saved')
